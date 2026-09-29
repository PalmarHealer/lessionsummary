"""Plain-text extraction from uploaded documents.

The output is read by an LLM, not a human, so structure matters more than
fidelity: page/slide headings let the summary cite "Folie 7", and tables
become Markdown so rows and columns survive the trip.
"""
from __future__ import annotations

import csv
import io
import logging
from pathlib import Path
from typing import Any

from .errors import JobError

log = logging.getLogger(__name__)

# Spreadsheets in school are mostly grade lists and data sets; a few hundred
# rows carry the shape, the rest only inflates the context window.
MAX_SHEET_ROWS = 200
MAX_SHEET_COLS = 30

TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".tex", ".rtf"}


def _md_table(rows: list[list[Any]]) -> str:
    rows = [[("" if c is None else str(c)).replace("|", "\\|").replace("\n", " ").strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(path: Path) -> dict[str, Any]:
    try:
        import pymupdf
    except ImportError:  # pymupdf < 1.24 only ships the legacy name
        import fitz as pymupdf

    try:
        doc = pymupdf.open(str(path))
    except Exception as exc:
        raise JobError(f"could not open PDF: {exc}") from exc
    parts: list[str] = []
    empty = 0
    with doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if not text:
                empty += 1
            parts.append(f"## Seite {i}\n\n{text}".rstrip())
        pages = doc.page_count
    # Scanned worksheets have no text layer. Say so instead of returning
    # headings only, so the summariser knows to look at the PDF itself.
    meta = {"pages_without_text": empty}
    if pages and empty == pages:
        meta["scanned"] = True
    return {"text": "\n\n".join(parts), "pages": pages, "meta": meta}


def _docx(path: Path) -> dict[str, Any]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    parts: list[str] = []
    # Iterate the body in document order; `d.paragraphs` and `d.tables`
    # separately would move every table to the end.
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            if style.startswith("Heading") or style.startswith("Überschrift"):
                level = "".join(ch for ch in style if ch.isdigit()) or "1"
                parts.append("#" * min(int(level) + 1, 6) + " " + text)
            elif style.startswith("List") or style.startswith("Liste"):
                parts.append(f"- {text}")
            else:
                parts.append(text)
        elif tag == "tbl":
            t = Table(child, d)
            parts.append(_md_table([[cell.text for cell in row.cells] for row in t.rows]))
    return {"text": "\n\n".join(p for p in parts if p), "pages": None, "meta": {}}


def _pptx(path: Path) -> dict[str, Any]:
    from pptx import Presentation

    prs = Presentation(str(path))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        lines = [f"## Folie {i}"]
        for shape in slide.shapes:
            if getattr(shape, "has_table", False) and shape.has_table:
                lines.append(_md_table([[c.text for c in r.cells] for r in shape.table.rows]))
            elif getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(r.text for r in para.runs).strip()
                    if text:
                        lines.append(("  " * para.level) + f"- {text}" if para.level else text)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                lines.append(f"Notizen: {notes}")
        parts.append("\n".join(lines))
    return {"text": "\n\n".join(parts), "pages": len(prs.slides), "meta": {}}


def _xlsx(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts: list[str] = []
    truncated: list[str] = []
    try:
        for ws in wb.worksheets:
            rows: list[list[Any]] = []
            for n, row in enumerate(ws.iter_rows(values_only=True)):
                if n >= MAX_SHEET_ROWS:
                    truncated.append(ws.title)
                    break
                rows.append(list(row[:MAX_SHEET_COLS]))
            table = _md_table(rows)
            if table:
                parts.append(f"## Tabelle {ws.title}\n\n{table}")
    finally:
        wb.close()
    return {"text": "\n\n".join(parts), "pages": None, "meta": {"truncated_sheets": truncated} if truncated else {}}


def _odf(path: Path, ext: str) -> dict[str, Any]:
    try:
        from odf import teletype
        from odf.draw import Page
        from odf.opendocument import load
        from odf.table import Table, TableCell, TableRow
        from odf.text import P
    except ImportError as exc:
        raise JobError("unsupported: odfpy not installed") from exc

    doc = load(str(path))
    parts: list[str] = []
    if ext == ".odp":
        for i, page in enumerate(doc.getElementsByType(Page), start=1):
            texts = [teletype.extractText(p).strip() for p in page.getElementsByType(P)]
            parts.append("\n".join([f"## Folie {i}"] + [t for t in texts if t]))
        return {"text": "\n\n".join(parts), "pages": len(parts), "meta": {}}
    if ext == ".ods":
        for t in doc.spreadsheet.getElementsByType(Table):
            rows = []
            for n, r in enumerate(t.getElementsByType(TableRow)):
                if n >= MAX_SHEET_ROWS:
                    break
                rows.append([teletype.extractText(c) for c in r.getElementsByType(TableCell)][:MAX_SHEET_COLS])
            table = _md_table(rows)
            if table:
                parts.append(f"## Tabelle {t.getAttribute('name')}\n\n{table}")
        return {"text": "\n\n".join(parts), "pages": None, "meta": {}}
    # .odt: top-level blocks in document order; lists and tables flatten to text.
    for node in doc.text.childNodes:
        name = getattr(node, "qname", (None, None))[1]
        text = teletype.extractText(node).strip()
        if not text:
            continue
        if name == "h":
            level = int(node.getAttribute("outlinelevel") or 1)
            parts.append("#" * min(level + 1, 6) + " " + text)
        else:
            parts.append(text)
    return {"text": "\n\n".join(parts), "pages": None, "meta": {}}


def _plain(path: Path, ext: str) -> dict[str, Any]:
    text = _decode(path.read_bytes())
    if ext in (".csv", ".tsv"):
        dialect = "excel-tab" if ext == ".tsv" else "excel"
        try:
            rows = list(csv.reader(io.StringIO(text), dialect=dialect))[: MAX_SHEET_ROWS + 1]
            return {"text": _md_table(rows), "pages": None, "meta": {}}
        except csv.Error:
            pass
    return {"text": text, "pages": None, "meta": {}}


_MIME_EXT = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/vnd.oasis.opendocument.text": ".odt",
    "application/vnd.oasis.opendocument.presentation": ".odp",
    "application/vnd.oasis.opendocument.spreadsheet": ".ods",
    "text/plain": ".txt",
    "text/markdown": ".md",
    "text/csv": ".csv",
}


def detect_ext(name: str, mime: str | None) -> str:
    ext = Path(name or "").suffix.lower()
    if ext:
        return ext
    return _MIME_EXT.get((mime or "").split(";")[0].strip().lower(), "")


def extract(path: Path, name: str, mime: str | None) -> dict[str, Any]:
    ext = detect_ext(name, mime)
    try:
        if ext == ".pdf":
            return _pdf(path)
        if ext in (".docx", ".docm"):
            return _docx(path)
        if ext in (".pptx", ".pptm"):
            return _pptx(path)
        if ext in (".xlsx", ".xlsm"):
            return _xlsx(path)
        if ext in (".odt", ".odp", ".ods"):
            return _odf(path, ext)
        if ext in TEXT_EXT or (mime or "").startswith("text/"):
            return _plain(path, ext)
    except JobError:
        raise
    except ImportError as exc:
        # A missing parser is a deployment problem, another worker may have it.
        raise JobError(f"extractor dependency missing: {exc}", retryable=True) from exc
    except Exception as exc:
        raise JobError(f"could not read {ext or 'file'}: {exc}") from exc
    # Legacy .doc/.ppt/.xls would need LibreOffice; not worth a 500 MB image.
    raise JobError(f"unsupported file type {ext or mime or 'unknown'}")
